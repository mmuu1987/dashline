from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOURCE_DIR = Path(__file__).resolve().parent
OUTPUT = SOURCE_DIR / "潮汐猎场-完整游戏策划案.docx"
SOURCES = [
    "README.md",
    "00-playtest-impressions.md",
    "01-product-and-core-loop.md",
    "02-gameplay-systems.md",
    "03-content-and-level-design.md",
    "04-progression-and-economy.md",
    "05-ui-art-audio.md",
    "06-balance-data-qa.md",
    "07-evidence-and-design-rationale.md",
    "08-monetization-and-liveops.md",
]


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"(`{1,2}|\*\*|__)(.+?)\1", r"\2", text)
    return text.replace("\\|", "|").strip()


def add_text_paragraph(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(clean_inline(text))
    set_east_asia_font(run, "微软雅黑")


def add_code_block(document: Document, language: str, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="策划案代码")
    if language:
        label = paragraph.add_run(f"[{language}]\n")
        label.bold = True
        label.font.color.rgb = RGBColor(73, 101, 125)
        set_east_asia_font(label, "Consolas")
    body = paragraph.add_run("\n".join(lines))
    set_east_asia_font(body, "Consolas")
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F6F8")
    p_pr.append(shd)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        rows.append([clean_inline(cell.strip()) for cell in raw.split("|")])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    return rows, i


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(columns):
            value = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_east_asia_font(run, "微软雅黑")
                    run.font.size = Pt(9)
                    if r_idx == 0:
                        run.bold = True
            if r_idx == 0:
                shade_cell(cell, "DCEAF3")


def add_markdown(document: Document, path: Path, skip_first_h1: bool = False) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_language = ""
    code_lines: list[str] = []
    first_h1_skipped = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(document, code_language, code_lines)
                in_code = False
                code_language = ""
                code_lines = []
            else:
                in_code = True
                code_language = stripped[3:].strip()
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            if skip_first_h1 and level == 1 and not first_h1_skipped:
                first_h1_skipped = True
                i += 1
                continue
            paragraph = document.add_heading(clean_inline(heading.group(2)), level=level)
            for run in paragraph.runs:
                set_east_asia_font(run, "微软雅黑")
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(document, rows)
            continue
        if stripped.startswith(">"):
            add_text_paragraph(document, stripped.lstrip("> "), "策划案引用")
            i += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            add_text_paragraph(document, bullet.group(1), "List Bullet")
            i += 1
            continue
        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if numbered:
            add_text_paragraph(document, numbered.group(1), "List Number")
            i += 1
            continue
        add_text_paragraph(document, stripped)
        i += 1
    if in_code:
        add_code_block(document, code_language, code_lines)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color in (
        ("Title", 30, "12344A"),
        ("Heading 1", 20, "0D5C63"),
        ("Heading 2", 15, "176B74"),
        ("Heading 3", 12, "2C7A7B"),
    ):
        style = document.styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    code = document.styles.add_style("策划案代码", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.right_indent = Cm(0.5)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    quote = document.styles.add_style("策划案引用", WD_STYLE_TYPE.PARAGRAPH)
    quote.font.name = "微软雅黑"
    quote._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    quote.font.size = Pt(10)
    quote.font.italic = True
    quote.font.color.rgb = RGBColor(83, 96, 106)
    quote.paragraph_format.left_indent = Cm(0.8)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE")


def build() -> Path:
    document = Document()
    configure_styles(document)
    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    add_page_number(section)

    document.add_paragraph()
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("《潮汐猎场》")
    title_run.bold = True
    title_run.font.size = Pt(34)
    title_run.font.color.rgb = RGBColor(13, 92, 99)
    set_east_asia_font(title_run, "微软雅黑")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("完整游戏策划案\n含《鱼吃鱼》试玩感想与同类产品设计")
    subtitle_run.font.size = Pt(17)
    subtitle_run.font.color.rgb = RGBColor(73, 101, 125)
    set_east_asia_font(subtitle_run, "微软雅黑")
    document.add_paragraph()
    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status_run = status.add_run("独立新项目提案 · 工作名 · 2026-09-02")
    status_run.font.size = Pt(11)
    set_east_asia_font(status_run, "微软雅黑")

    document.add_page_break()
    document.add_heading("目录", level=1)
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    hint = document.add_paragraph("如目录未自动显示，请在 Word 中右键目录并选择“更新域”。")
    for run in hint.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(120, 120, 120)
        set_east_asia_font(run, "微软雅黑")

    document.add_page_break()
    for index, source in enumerate(SOURCES):
        if index > 0:
            document.add_section(WD_SECTION.NEW_PAGE)
        add_markdown(document, SOURCE_DIR / source, skip_first_h1=index == 0)

    props = document.core_properties
    props.title = "《潮汐猎场》完整游戏策划案"
    props.subject = "《鱼吃鱼》试玩感想与同类微信小游戏策划案"
    props.author = "项目策划组"
    props.keywords = "游戏策划, 大鱼吃小鱼, 微信小游戏, 潮汐猎场"
    document.save(OUTPUT)
    return OUTPUT


def build_supplement(source: str, output_name: str, title_text: str, subtitle_text: str) -> Path:
    document = Document()
    configure_styles(document)
    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    add_page_number(section)

    document.add_paragraph()
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(title_text)
    title_run.bold = True
    title_run.font.size = Pt(30)
    title_run.font.color.rgb = RGBColor(13, 92, 99)
    set_east_asia_font(title_run, "微软雅黑")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(subtitle_text)
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor(73, 101, 125)
    set_east_asia_font(subtitle_run, "微软雅黑")
    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status_run = status.add_run("首版执行基线 · 2026-09-02")
    status_run.font.size = Pt(11)
    set_east_asia_font(status_run, "微软雅黑")

    document.add_page_break()
    document.add_heading("目录", level=1)
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    hint = document.add_paragraph("如目录未自动显示，请在 Word 中右键目录并选择“更新域”。")
    for run in hint.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(120, 120, 120)
        set_east_asia_font(run, "微软雅黑")

    document.add_page_break()
    add_markdown(document, SOURCE_DIR / source, skip_first_h1=True)
    output = SOURCE_DIR / output_name
    props = document.core_properties
    props.title = title_text
    props.subject = subtitle_text
    props.author = "项目策划组"
    props.keywords = "游戏开发, 执行流程, 资产清单, 微信小游戏, 潮汐猎场"
    document.save(output)
    return output


if __name__ == "__main__":
    outputs: list[Path] = []
    try:
        outputs.append(build())
    except PermissionError:
        print(f"skip locked file: {OUTPUT}")
    outputs.extend([
        build_supplement(
            "09-development-execution-plan.md",
            "潮汐猎场-开发执行流程.docx",
            "《潮汐猎场》开发执行流程",
            "16 周首版开发计划 · 团队分工 · 关键依赖 · 验收与发布",
        ),
        build_supplement(
            "10-asset-specification-checklist.md",
            "潮汐猎场-资产设定清单.docx",
            "《潮汐猎场》资产设定清单",
            "角色 · 场景 · 动画 · 特效 · UI · 音频 · 技术规格与验收",
        ),
        build_supplement(
            "11-minimum-mvp.md",
            "潮汐猎场-最小MVP方案.docx",
            "《潮汐猎场》最小 MVP 方案",
            "4 周核心乐趣验证 · 最小范围 · 资产与验收门槛",
        ),
    ])
    for output in outputs:
        print(output)
